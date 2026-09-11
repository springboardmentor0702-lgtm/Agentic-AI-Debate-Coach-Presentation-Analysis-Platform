import io
from xml.sax.saxutils import escape


def build_pdf(report: dict) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, title="Debate Analysis Report")
    styles = getSampleStyleSheet()
    story = []
    story.append(Paragraph("Debate Coach && Presentation Analysis Report", styles["Title"]))
    story.append(Paragraph(f"Topic: {escape(str(report.get('topic', '')))}", styles["Normal"]))
    story.append(Paragraph(f"Generated: {escape(str(report.get('generated_at', '')))}", styles["Normal"]))
    story.append(Spacer(1, 8 * mm))

    s = report.get("scores", {})
    story.append(Paragraph(f"Overall Weighted Score: <b>{s.get('overall_score', 0)}/100</b>", styles["Heading2"]))
    rows = [["Component", "Weight", "Score"],
            ["Argument Quality", "30%", str(s.get("argument_quality", ""))],
            ["Evidence Usage", "20%", str(s.get("evidence_usage", ""))],
            ["Logical Consistency", "20%", str(s.get("logical_consistency", ""))],
            ["Rebuttal Effectiveness", "15%", str(s.get("rebuttal_effectiveness", ""))],
            ["Communication Skills", "15%", str(s.get("communication_skills", ""))]]
    t = Table(rows, colWidths=[60 * mm, 25 * mm, 25 * mm])
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey)]))
    story.append(t)
    story.append(Spacer(1, 6 * mm))

    pm = report.get("presentation_metrics") or {}
    if pm:
        story.append(Paragraph("Presentation Metrics", styles["Heading2"]))
        pt = Table([[k.replace("_", " ").title(), str(v)] for k, v in pm.items()],
                   colWidths=[60 * mm, 50 * mm])
        pt.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey)]))
        story.append(pt)
        story.append(Spacer(1, 6 * mm))

    ff = report.get("fallacies_found", [])
    story.append(Paragraph(f"Logical Fallacies Detected: {len(ff)}", styles["Heading2"]))
    for f in ff:
        story.append(Paragraph(
            f"<b>{escape(str(f.get('fallacy', '')).replace('_', ' ').title())}</b>: "
            f"{escape(str(f.get('explanation', '')))} Fix: {escape(str(f.get('correction', '')))}",
            styles["Normal"]))
        story.append(Spacer(1, 2 * mm))

    ch = report.get("coaching", {})
    story.append(Paragraph("Personalized Coaching", styles["Heading2"]))
    story.append(Paragraph(escape(str(ch.get("coaching_feedback", ""))), styles["Normal"]))
    for r in ch.get("improvement_recommendations", []):
        story.append(Paragraph(f"- {escape(str(r))}", styles["Normal"]))
    story.append(Spacer(1, 6 * mm))

    story.append(Paragraph("Transcript", styles["Heading2"]))
    for t2 in report.get("transcript", []):
        who = "You" if t2.get("speaker") == "user" else "AI Opponent"
        story.append(Paragraph(f"<b>{who}:</b> {escape(str(t2.get('content', '')))}", styles["Normal"]))
        story.append(Spacer(1, 2 * mm))

    doc.build(story)
    return buf.getvalue()


def build_excel(report: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    s = report.get("scores", {})

    ws = wb.active
    ws.title = "Summary"
    ws.append(["Debate Analysis Report"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Topic", report.get("topic", "")])
    ws.append(["Generated", str(report.get("generated_at", ""))])
    ws.append([])
    ws.append(["Overall Score", s.get("overall_score")])

    ws2 = wb.create_sheet("Scores")
    ws2.append(["Component", "Score"])
    for k in ["argument_quality", "evidence_usage", "logical_consistency",
              "rebuttal_effectiveness", "communication_skills"]:
        ws2.append([k.replace("_", " ").title(), s.get(k)])

    ws3 = wb.create_sheet("Presentation")
    ws3.append(["Metric", "Value"])
    for k, v in (report.get("presentation_metrics") or {}).items():
        ws3.append([k.replace("_", " ").title(), v])

    ws4 = wb.create_sheet("Fallacies")
    ws4.append(["Fallacy", "Explanation", "Correction", "Confidence"])
    for f in report.get("fallacies_found", []):
        ws4.append([f.get("fallacy"), f.get("explanation"), f.get("correction"), f.get("confidence")])

    ws5 = wb.create_sheet("Transcript")
    ws5.append(["Speaker", "Content"])
    for t in report.get("transcript", []):
        ws5.append(["You" if t.get("speaker") == "user" else "AI", t.get("content", "")])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
