"""Presentation management endpoints."""

from datetime import datetime, timezone
from pathlib import Path
import re
from typing import Annotated

from fastapi import (
    APIRouter,
    Depends,
    File,
    Form,
    HTTPException,
    UploadFile,
    status,
)
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.dependencies import get_current_user, get_db
from app.core.settings import settings
from app.models.models import (
    Presentation,
    PresentationAnalysis,
    User,
)
from app.services.gemini import analyze_presentation_text


router = APIRouter(
    prefix="/presentations",
    tags=["presentations"],
)


# ============================================================================
# Schemas
# ============================================================================


class PresentationAnalysisSchema(BaseModel):
    """Presentation analysis response schema."""

    id: int
    presentation_id: int

    transcript: str | None = None

    pace: float | None = None

    filler_word_usage: int | None = None

    clarity_score: float | None = None

    confidence_score: float | None = None

    engagement_score: float | None = None

    speaking_duration_seconds: float | None = None

    words_per_minute: float | None = Nonertt

    long_pauses: int | None = None

    ai_feedback: str | None = None

    created_at: datetime

    model_config = {
        "from_attributes": True
    }


class PresentationSchema(BaseModel):
    """Presentation response schema."""

    id: int
    user_id: int

    title: str

    file_name: str

    file_path: str

    status: str

    created_at: datetime

    updated_at: datetime

    analysis: PresentationAnalysisSchema | None = None

    model_config = {
        "from_attributes": True
    }


class CreatePresentationRequest(BaseModel):
    """Create presentation request schema."""

    title: str = Field(
        ...,
        min_length=3,
        max_length=255,
    )

    file_name: str = Field(
        ...,
        min_length=1,
        max_length=255,
    )

    file_path: str = Field(
        ...,
        min_length=1,
        max_length=500,
    )


class UpdatePresentationRequest(BaseModel):
    """Update presentation request schema."""

    title: str | None = Field(
        None,
        min_length=3,
        max_length=255,
    )

    file_name: str | None = Field(
        None,
        min_length=1,
        max_length=255,
    )

    file_path: str | None = Field(
        None,
        min_length=1,
        max_length=500,
    )


class AnalyzePresentationRequest(BaseModel):
    """Presentation analysis request schema."""

    transcript: str | None = None

    pace: float | None = None

    clarity_score: float | None = Field(
        None,
        ge=0,
        le=100,
    )

    confidence_score: float | None = Field(
        None,
        ge=0,
        le=100,
    )

    engagement_score: float | None = Field(
        None,
        ge=0,
        le=100,
    )

    ai_feedback: str | None = None


# ============================================================================
# Create / Upload Presentation
# ============================================================================


@router.post(
    "",
    response_model=PresentationSchema,
    status_code=status.HTTP_201_CREATED,
)
def create_presentation(
    current_user: Annotated[
        User,
        Depends(get_current_user),
    ],
    db: Annotated[
        Session,
        Depends(get_db),
    ],
    title: str = Form(
        ...,
        min_length=3,
        max_length=255,
    ),
    file: UploadFile = File(...),
) -> PresentationSchema:
    """
    Upload and persist a PDF or PPTX presentation.
    """

    if not file.filename:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="File name is required.",
        )

    suffix = Path(file.filename).suffix.lower()

    if suffix not in {".pdf", ".pptx"}:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Only PDF and PPTX files are supported.",
        )

    try:
        content = file.file.read()
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Unable to read uploaded file: {str(exc)}",
        ) from exc

    max_size = (
        settings.max_upload_size_mb
        * 1024
        * 1024
    )

    if len(content) > max_size:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail=(
                f"Presentation file is too large. "
                f"Maximum size is {settings.max_upload_size_mb} MB."
            ),
        )

    upload_dir = Path(
        settings.upload_directory
    )

    upload_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    safe_name = re.sub(
        r"[^A-Za-z0-9_.-]",
        "_",
        file.filename,
    )

    stored_path = upload_dir / safe_name

    try:
        stored_path.write_bytes(content)
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Unable to save presentation: {str(exc)}",
        ) from exc

    presentation = Presentation(
        user_id=current_user.id,
        title=title,
        file_name=file.filename,
        file_path=str(stored_path),
        status="UPLOADED",
    )

    db.add(presentation)
    db.commit()
    db.refresh(presentation)

    return PresentationSchema.model_validate(
        presentation
    )


# ============================================================================
# List Presentations
# ============================================================================


@router.get(
    "",
    response_model=list[PresentationSchema],
)
def list_presentations(
    current_user: Annotated[
        User,
        Depends(get_current_user),
    ],
    db: Annotated[
        Session,
        Depends(get_db),
    ],
) -> list[PresentationSchema]:
    """List all presentations for the current user."""

    stmt = (
        select(Presentation)
        .where(
            Presentation.user_id == current_user.id
        )
        .order_by(
            Presentation.created_at.desc()
        )
    )

    result = db.execute(stmt)

    presentations = result.scalars().all()

    return [
        PresentationSchema.model_validate(p)
        for p in presentations
    ]


# ============================================================================
# Get Presentation
# ============================================================================


@router.get(
    "/{presentation_id}",
    response_model=PresentationSchema,
)
def get_presentation(
    presentation_id: int,
    current_user: Annotated[
        User,
        Depends(get_current_user),
    ],
    db: Annotated[
        Session,
        Depends(get_db),
    ],
) -> PresentationSchema:
    """Get a specific presentation."""

    stmt = select(Presentation).where(
        (Presentation.id == presentation_id)
        & (
            Presentation.user_id
            == current_user.id
        )
    )

    result = db.execute(stmt)

    presentation = result.scalar_one_or_none()

    if not presentation:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Presentation not found.",
        )

    return PresentationSchema.model_validate(
        presentation
    )


# ============================================================================
# Update Presentation
# ============================================================================


@router.put(
    "/{presentation_id}",
    response_model=PresentationSchema,
)
def update_presentation(
    presentation_id: int,
    request: UpdatePresentationRequest,
    current_user: Annotated[
        User,
        Depends(get_current_user),
    ],
    db: Annotated[
        Session,
        Depends(get_db),
    ],
) -> PresentationSchema:
    """Update a presentation."""

    stmt = select(Presentation).where(
        (Presentation.id == presentation_id)
        & (
            Presentation.user_id
            == current_user.id
        )
    )

    result = db.execute(stmt)

    presentation = result.scalar_one_or_none()

    if not presentation:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Presentation not found.",
        )

    if request.title is not None:
        presentation.title = request.title

    if request.file_name is not None:
        presentation.file_name = request.file_name

    if request.file_path is not None:
        presentation.file_path = request.file_path

    presentation.updated_at = datetime.now(
        timezone.utc
    )

    db.commit()
    db.refresh(presentation)

    return PresentationSchema.model_validate(
        presentation
    )


# ============================================================================
# Delete Presentation
# ============================================================================


@router.delete(
    "/{presentation_id}",
    status_code=status.HTTP_204_NO_CONTENT,
)
def delete_presentation(
    presentation_id: int,
    current_user: Annotated[
        User,
        Depends(get_current_user),
    ],
    db: Annotated[
        Session,
        Depends(get_db),
    ],
) -> None:
    """Delete a presentation."""

    stmt = select(Presentation).where(
        (Presentation.id == presentation_id)
        & (
            Presentation.user_id
            == current_user.id
        )
    )

    result = db.execute(stmt)

    presentation = result.scalar_one_or_none()

    if not presentation:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Presentation not found.",
        )

    # Delete the database record.
    db.delete(presentation)
    db.commit()


# ============================================================================
# Analyze Presentation With Gemini
# ============================================================================


@router.post(
    "/{presentation_id}/analyze",
    response_model=PresentationAnalysisSchema,
    status_code=status.HTTP_201_CREATED,
)
def analyze_presentation(
    presentation_id: int,
    current_user: Annotated[
        User,
        Depends(get_current_user),
    ],
    db: Annotated[
        Session,
        Depends(get_db),
    ],
) -> PresentationAnalysisSchema:
    """
    Extract presentation text and analyze it using Gemini.
    """

    # ------------------------------------------------------------------------
    # Verify presentation exists and belongs to current user
    # ------------------------------------------------------------------------

    stmt = select(Presentation).where(
        (Presentation.id == presentation_id)
        & (
            Presentation.user_id
            == current_user.id
        )
    )

    result = db.execute(stmt)

    presentation = result.scalar_one_or_none()

    if not presentation:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Presentation not found.",
        )

    # ------------------------------------------------------------------------
    # Check physical file
    # ------------------------------------------------------------------------

    file_path = Path(
        presentation.file_path
    )

    if not file_path.exists():
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=(
                "Presentation file was not found "
                "on the server."
            ),
        )

    # ------------------------------------------------------------------------
    # Extract PDF/PPTX text
    # ------------------------------------------------------------------------

    try:
        extracted_text = extract_presentation_text(
            presentation.file_path,
            presentation.file_name,
        )
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=(
                f"Failed to extract presentation text: "
                f"{str(exc)}"
            ),
        ) from exc

    if not extracted_text.strip():
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                "No readable text was found in "
                "the presentation."
            ),
        )

    # ------------------------------------------------------------------------
    # Analyze extracted content with Gemini
    # ------------------------------------------------------------------------

    try:
        ai_result = analyze_presentation_text(
            extracted_text,
            presentation.title,
        )
    except ValueError as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=str(exc),
        ) from exc
    except RuntimeError as exc:
        raise HTTPException(
            status_code=status.HTTP_502_BAD_GATEWAY,
            detail=str(exc),
        ) from exc
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=(
                f"Presentation analysis failed: "
                f"{str(exc)}"
            ),
        ) from exc

    # ------------------------------------------------------------------------
    # Check if analysis already exists
    # ------------------------------------------------------------------------

    stmt = select(PresentationAnalysis).where(
        PresentationAnalysis.presentation_id
        == presentation_id
    )

    result = db.execute(stmt)

    analysis = result.scalar_one_or_none()

    # ------------------------------------------------------------------------
    # Update existing analysis
    # ------------------------------------------------------------------------

    if analysis:

        analysis.transcript = ai_result.get(
            "transcript"
        )

        analysis.pace = ai_result.get(
            "pace"
        )

        analysis.clarity_score = ai_result.get(
            "clarity_score"
        )

        analysis.confidence_score = ai_result.get(
            "confidence_score"
        )

        analysis.engagement_score = ai_result.get(
            "engagement_score"
        )

        analysis.ai_feedback = ai_result.get(
            "ai_feedback"
        )

    # ------------------------------------------------------------------------
    # Create new analysis
    # ------------------------------------------------------------------------

    else:

        analysis = PresentationAnalysis(
            presentation_id=presentation_id,
            transcript=ai_result.get(
                "transcript"
            ),
            pace=ai_result.get(
                "pace"
            ),
            clarity_score=ai_result.get(
                "clarity_score"
            ),
            confidence_score=ai_result.get(
                "confidence_score"
            ),
            engagement_score=ai_result.get(
                "engagement_score"
            ),
            ai_feedback=ai_result.get(
                "ai_feedback"
            ),

            # These require audio/speech analysis.
            filler_word_usage=None,
            speaking_duration_seconds=None,
            words_per_minute=None,
            long_pauses=None,
        )

        db.add(analysis)

    # ------------------------------------------------------------------------
    # Update presentation status
    # ------------------------------------------------------------------------

    presentation.status = "ANALYZED"

    presentation.updated_at = datetime.now(
        timezone.utc
    )

    # ------------------------------------------------------------------------
    # Save
    # ------------------------------------------------------------------------

    try:
        db.commit()
        db.refresh(analysis)
    except Exception as exc:
        db.rollback()

        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=(
                f"Failed to save presentation analysis: "
                f"{str(exc)}"
            ),
        ) from exc

    return PresentationAnalysisSchema.model_validate(
        analysis
    )


# ============================================================================
# Extract Presentation Text
# ============================================================================


def extract_presentation_text(
    file_path: str,
    file_name: str,
) -> str:
    """
    Extract text from PDF or PPTX.
    """

    suffix = Path(
        file_name
    ).suffix.lower()

    # ------------------------------------------------------------------------
    # PDF
    # ------------------------------------------------------------------------

    if suffix == ".pdf":

        try:
            from pypdf import PdfReader

            reader = PdfReader(
                file_path
            )

            pages = []

            for page in reader.pages:

                page_text = (
                    page.extract_text()
                    or ""
                )

                if page_text.strip():
                    pages.append(
                        page_text
                    )

            return "\n\n".join(
                pages
            )[:50000]

        except Exception as exc:
            raise RuntimeError(
                f"Failed to read PDF: {str(exc)}"
            ) from exc

    # ------------------------------------------------------------------------
    # PPTX
    # ------------------------------------------------------------------------

    if suffix == ".pptx":

        try:
            from pptx import Presentation as PptxPresentation

            deck = PptxPresentation(
                file_path
            )

            slides_text = []

            for slide_number, slide in enumerate(
                deck.slides,
                start=1,
            ):

                slide_parts = []

                for shape in slide.shapes:

                    if hasattr(
                        shape,
                        "text",
                    ):

                        text = (
                            shape.text
                            or ""
                        ).strip()

                        if text:
                            slide_parts.append(
                                text
                            )

                if slide_parts:

                    slides_text.append(
                        f"Slide {slide_number}\n"
                        + "\n".join(
                            slide_parts
                        )
                    )

            return "\n\n".join(
                slides_text
            )[:50000]

        except Exception as exc:
            raise RuntimeError(
                f"Failed to read PPTX: {str(exc)}"
            ) from exc

    raise HTTPException(
        status_code=status.HTTP_400_BAD_REQUEST,
        detail="Unsupported presentation format.",
    )


# ============================================================================
# Get Presentation Analysis
# ============================================================================


@router.get(
    "/{presentation_id}/analysis",
    response_model=PresentationAnalysisSchema,
)
def get_presentation_analysis(
    presentation_id: int,
    current_user: Annotated[
        User,
        Depends(get_current_user),
    ],
    db: Annotated[
        Session,
        Depends(get_db),
    ],
) -> PresentationAnalysisSchema:
    """Get analysis for a specific presentation."""

    # ------------------------------------------------------------------------
    # Verify presentation belongs to current user
    # ------------------------------------------------------------------------

    stmt = select(Presentation).where(
        (Presentation.id == presentation_id)
        & (
            Presentation.user_id
            == current_user.id
        )
    )

    result = db.execute(stmt)

    presentation = result.scalar_one_or_none()

    if not presentation:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Presentation not found.",
        )

    # ------------------------------------------------------------------------
    # Get analysis
    # ------------------------------------------------------------------------

    stmt = select(
        PresentationAnalysis
    ).where(
        PresentationAnalysis.presentation_id
        == presentation_id
    )

    result = db.execute(stmt)

    analysis = result.scalar_one_or_none()

    if not analysis:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Analysis not found.",
        )

    return PresentationAnalysisSchema.model_validate(
        analysis
    )